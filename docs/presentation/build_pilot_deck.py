"""Generate a client-facing PowerPoint for Test Automation Pilot.

Design language mirrors raptor_framework_slides.pptx: 13.33x7.5 widescreen,
white background, Ericsson Hilda / Ericsson Hilda Light type, a dark title
with a grey divider rule, coloured accent chips, numbered step badges, and a
thin footer rule with a page number. Content is drawn from the tool's real
capabilities (MCP tools in mcp_server/server.py, the five Streamlit features
in _pages/home.py, and the LLM/MCP wiring in docs/MCP_SERVER_SETUP.md).

Run:  python docs/presentation/build_pilot_deck.py
Output: docs/presentation/Test_Automation_Pilot.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ── Palette (matched to the reference deck) ─────────────────────────
INK = RGBColor(0x18, 0x18, 0x18)      # near-black title/body-strong
GREY = RGBColor(0x55, 0x55, 0x55)     # secondary text
RULE = RGBColor(0xDD, 0xDD, 0xDD)     # title divider
FOOT_RULE = RGBColor(0xE2, 0xE2, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BLUE = RGBColor(0x11, 0x74, 0xE6)
BLUE_DK = RGBColor(0x0B, 0x4E, 0xA8)
BLUE_CHIP = RGBColor(0xEA, 0xF1, 0xFD)
BLUE_FAINT = RGBColor(0xFA, 0xFC, 0xFF)

GREEN = RGBColor(0x15, 0x8A, 0x4D)
GREEN_CHIP = RGBColor(0xE8, 0xF7, 0xEE)
RED = RGBColor(0xE6, 0x5D, 0x6A)
RED_CHIP = RGBColor(0xFC, 0xEC, 0xEE)
ORANGE = RGBColor(0xE6, 0x6E, 0x19)
ORANGE_CHIP = RGBColor(0xFD, 0xF2, 0xDF)
PURPLE = RGBColor(0x6B, 0x4E, 0xE6)
PURPLE_CHIP = RGBColor(0xEF, 0xEB, 0xFD)

CODE_BG = RGBColor(0x1E, 0x24, 0x2B)
CODE_FG = RGBColor(0xDC, 0xE6, 0xF2)

FONT = "Ericsson Hilda"
FONT_LIGHT = "Ericsson Hilda Light"
MONO = "Consolas"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _no_autofit(tf):
    tf.word_wrap = True


def _set_run(run, *, size, color, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def textbox(slide, l, t, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    _no_autofit(tf)
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
    return box, tf


def para(tf, text, *, size, color, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         space_after=4, space_before=0, first=False, line_spacing=None):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, color=color, bold=bold, font=font)
    return p


def rounded(slide, l, t, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    _solid(shp, color)
    # soften the corner radius
    try:
        shp.adjustments[0] = 0.08
    except Exception:
        pass
    return shp


def rect(slide, l, t, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    _solid(shp, color)
    return shp


def _page_chrome(slide, title, page_no, total, kicker="TEST AUTOMATION PILOT"):
    # kicker
    _, ktf = textbox(slide, 0.60, 0.33, 10.5, 0.28)
    para(ktf, kicker, size=10, color=BLUE, bold=True, first=True)
    # title
    _, ttf = textbox(slide, 0.60, 0.62, 12.1, 0.6)
    para(ttf, title, size=24, color=INK, bold=True, font=FONT_LIGHT, first=True)
    # divider
    rect(slide, 0.60, 1.28, 12.13, 0.02, RULE)
    # footer rule + page number
    rect(slide, 0.60, 7.05, 12.13, 0.012, FOOT_RULE)
    _, ftf = textbox(slide, 11.23, 7.13, 1.50, 0.28)
    para(ftf, f"{page_no:02d} / {total:02d}", size=9, color=GREY, align=PP_ALIGN.RIGHT, first=True)


def _takeaway(slide, text):
    bar = rounded(slide, 0.60, 6.40, 12.13, 0.5, BLUE_FAINT)
    bar.line.color.rgb = BLUE_CHIP
    bar.line.width = Pt(0.75)
    _, tf = textbox(slide, 0.80, 6.40, 11.7, 0.5, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, "", size=12.5, color=INK, first=True)
    r = p.add_run(); _set_run(r, size=12.5, color=BLUE, bold=True); r.text = "Takeaway  "
    r2 = p.add_run(); _set_run(r2, size=12, color=INK, bold=False); r2.text = text
    p.alignment = PP_ALIGN.CENTER


def chip_header(slide, l, t, w, text, chip_color, text_color):
    bar = rounded(slide, l, t, w, 0.40, chip_color)
    _, tf = textbox(slide, l + 0.15, t, w - 0.3, 0.40, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, size=12, color=text_color, bold=True, align=PP_ALIGN.CENTER, first=True)


def bullets(slide, l, t, w, h, items, color, size=11.5, bullet="•  ", gap=5):
    _, tf = textbox(slide, l, t, w, h)
    for i, it in enumerate(items):
        para(tf, bullet + it, size=size, color=color, first=(i == 0),
             space_after=gap, line_spacing=1.03)
    return tf


# ════════════════════════════════════════════════════════════════════
def build():
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    blank = prs.slide_layouts[6]
    TOTAL = 5

    # ─────────────────────────── SLIDE 1 — TITLE ───────────────────────
    s = prs.slides.add_slide(blank)
    # left accent band
    rect(s, 0.0, 0.0, 0.22, 7.5, BLUE)
    _, ktf = textbox(s, 0.9, 2.05, 11, 0.4)
    para(ktf, "AI-POWERED TESTING ADVISOR", size=14, color=BLUE, bold=True, first=True)
    _, ttf = textbox(s, 0.88, 2.5, 11.5, 1.4)
    para(ttf, "Test Automation Pilot", size=52, color=INK, bold=True, font=FONT_LIGHT, first=True)
    _, stf = textbox(s, 0.9, 3.95, 11.3, 1.1)
    para(stf, "Evaluate, migrate, and verify test automation frameworks — "
              "an AI assistant that recommends the right framework, converts "
              "test suites between frameworks, generates automated tests from "
              "manual cases, and proves nothing was lost in migration.",
         size=15, color=GREY, first=True, line_spacing=1.15)
    # bottom pills
    px = 0.9
    for label, col, chip in [("Framework Advisor", BLUE, BLUE_CHIP),
                             ("Code Migration", GREEN, GREEN_CHIP),
                             ("Test Generation", ORANGE, ORANGE_CHIP),
                             ("Coverage & Parity", PURPLE, PURPLE_CHIP)]:
        w = 0.28 + len(label) * 0.095
        chip_header(s, px, 5.5, w, label, chip, col)
        px += w + 0.25
    rect(s, 0.9, 6.45, 11.5, 0.012, FOOT_RULE)
    _, ftf = textbox(s, 0.9, 6.6, 11.5, 0.3)
    para(ftf, "Client Briefing  •  Capabilities, Architecture & Features",
         size=11, color=GREY, first=True)

    # ─────────────────── SLIDE 2 — WHAT IT DOES / WHY USEFUL ────────────
    s = prs.slides.add_slide(blank)
    _page_chrome(s, "What Test Automation Pilot Does", 2, TOTAL)
    _, itf = textbox(s, 0.60, 1.42, 12.13, 0.5)
    para(itf, "One assistant across the full framework lifecycle — from "
              "picking a framework, to migrating existing suites, to proving "
              "the migration is complete. It runs as a web app and as an MCP "
              "server any AI client can call.",
         size=12.5, color=GREY, first=True, line_spacing=1.1)

    cards = [
        ("CHOOSE", BLUE, BLUE_CHIP,
         "Recommend & score frameworks",
         ["Weighted scorecards across 7 criteria",
          "Compare any frameworks side by side",
          "20+ frameworks in a curated knowledge base"]),
        ("MIGRATE", GREEN, GREEN_CHIP,
         "Convert test code between frameworks",
         ["Selenium <-> Playwright, K6 -> Robot, etc.",
          "Every assertion preserved & auto-validated",
          "Phased migration roadmaps with effort sizing"]),
        ("GENERATE", ORANGE, ORANGE_CHIP,
         "Manual test cases -> runnable code",
         ["Playwright, Selenium, Cypress, Robot...",
          "Page objects, fixtures, run commands",
          "Boilerplate project scaffolding"]),
        ("VERIFY", PURPLE, PURPLE_CHIP,
         "Prove coverage & parity",
         ["Test-case x framework coverage matrix",
          "Cross-repo assertion-parity reports",
          "Catches silently dropped tests"]),
    ]
    cw, gap = 2.92, 0.17
    x = 0.60
    for tag, col, chip, title, items in cards:
        card = rounded(s, x, 2.15, cw, 3.78, BLUE_FAINT)
        card.line.color.rgb = chip
        card.line.width = Pt(1)
        chip_header(s, x + 0.15, 2.32, cw - 0.3, tag, chip, col)
        _, ttf = textbox(s, x + 0.2, 2.92, cw - 0.4, 0.75)
        para(ttf, title, size=13, color=INK, bold=True, first=True, line_spacing=1.05)
        bullets(s, x + 0.2, 3.78, cw - 0.35, 2.0, items, GREY, size=11, gap=10)
        x += cw + gap

    _takeaway(s, "Replaces scattered manual judgement, hand-porting, and "
                 "guesswork with one AI-assisted, evidence-backed workflow.")

    # ─────────────────── SLIDE 3 — ARCHITECTURE ────────────────────────
    s = prs.slides.add_slide(blank)
    _page_chrome(s, "Architecture — How It Fits Together", 3, TOTAL)

    # Left column: layered boxes
    def layer(l, t, w, h, title, sub, chip, col, title_sz=11.5):
        box = rounded(s, l, t, w, h, chip)
        box.line.color.rgb = col
        box.line.width = Pt(1)
        _, tf = textbox(s, l + 0.15, t, w - 0.3, h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, title, size=title_sz, color=col, bold=True, align=PP_ALIGN.CENTER, first=True)
        if sub:
            para(tf, sub, size=8.5, color=GREY, align=PP_ALIGN.CENTER, space_before=1)
        return box

    colL = 0.60
    colW = 6.0
    layer(colL, 1.5, colW, 0.62, "Users & AI Clients",
          "Testers / QA leads  •  Kiro & other MCP-aware assistants", BLUE_CHIP, BLUE_DK)
    conn = textbox(s, colL, 2.14, colW, 0.2)[1]
    para(conn, "v   uses", size=8, color=GREY, align=PP_ALIGN.CENTER, first=True)

    # two access surfaces side by side
    layer(colL, 2.36, colW / 2 - 0.08, 0.62, "Streamlit Web App",
          "5 guided feature pages", GREEN_CHIP, GREEN)
    layer(colL + colW / 2 + 0.08, 2.36, colW / 2 - 0.08, 0.62, "MCP Server",
          "20+ tools over stdio", PURPLE_CHIP, PURPLE)
    conn = textbox(s, colL, 3.0, colW, 0.2)[1]
    para(conn, "v   drive the same core", size=8, color=GREY, align=PP_ALIGN.CENTER, first=True)

    layer(colL, 3.22, colW, 0.9, "Core Engine",
          "Scoring  •  Migration planner  •  Code converter  •  Codegen  •  "
          "Coverage & repo-parity  •  Team profiles", BLUE_FAINT, BLUE, title_sz=12)
    conn = textbox(s, colL, 4.14, colW, 0.2)[1]
    para(conn, "v   grounded by / calls out to", size=8, color=GREY, align=PP_ALIGN.CENTER, first=True)

    layer(colL, 4.36, colW / 2 - 0.08, 0.95, "Knowledge Base",
          "Curated YAML + knowledge graph; 20+ frameworks; live discovery",
          ORANGE_CHIP, ORANGE)
    layer(colL + colW / 2 + 0.08, 4.36, colW / 2 - 0.08, 0.95, "LLM Provider",
          "Groq / Gemini API, or the AI client's own model via MCP sampling",
          RED_CHIP, RED)

    # LLM requirement callout
    note = rounded(s, colL, 5.55, colW, 0.78, RGBColor(0xFF, 0xF8, 0xE6))
    note.line.color.rgb = ORANGE
    note.line.width = Pt(1)
    _, ntf = textbox(s, colL + 0.18, 5.6, colW - 0.36, 0.7, anchor=MSO_ANCHOR.MIDDLE)
    p = para(ntf, "", size=10, color=INK, first=True)
    r = p.add_run(); _set_run(r, size=10.5, color=ORANGE, bold=True); r.text = "Needs an LLM:  "
    r2 = p.add_run(); _set_run(r2, size=10.5, color=INK)
    r2.text = ("AI features (chat, conversion, generation) connect to an LLM. "
               "Add a Groq or Gemini key in Settings — or, via MCP, the tool can "
               "use your AI client's own model (sampling), so no server key is spent.")

    # Right column: numbered request-flow steps
    steps = [
        ("User asks or triggers a feature",
         "In the web app, or from an AI client as an MCP tool call."),
        ("Core engine does the work",
         "Scoring, planning, conversion, generation, or parity checks."),
        ("Knowledge base grounds the answer",
         "Framework facts come from curated data, not guesses."),
        ("LLM fills the creative gaps",
         "Converts/generates code; everything is then validated."),
        ("Verified result returned",
         "Scorecards, converted code, parity reports — reproducible."),
    ]
    rx = 7.0
    ry = 1.5
    _, htf = textbox(s, rx, ry, 5.7, 0.3)
    para(htf, "REQUEST FLOW", size=10.5, color=BLUE, bold=True, first=True)
    ry += 0.42
    for i, (title, desc) in enumerate(steps, 1):
        badge = slide_badge = rounded(s, rx, ry, 0.34, 0.34, BLUE)
        badge.adjustments[0] = 0.5
        _, btf = textbox(s, rx, ry, 0.34, 0.34, anchor=MSO_ANCHOR.MIDDLE)
        para(btf, str(i), size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
        _, stf = textbox(s, rx + 0.5, ry - 0.04, 5.2, 0.28)
        para(stf, title, size=11.5, color=INK, bold=True, first=True)
        _, dtf = textbox(s, rx + 0.5, ry + 0.24, 5.2, 0.5)
        para(dtf, desc, size=9, color=GREY, first=True, line_spacing=1.05)
        ry += 0.86

    _takeaway(s, "Same core engine, two front doors (web app + MCP). An LLM "
                 "connection powers the AI features; the knowledge base keeps "
                 "answers grounded.")

    # ─────────────────── SLIDE 4 — FEATURES ────────────────────────────
    s = prs.slides.add_slide(blank)
    _page_chrome(s, "Five Core Features", 4, TOTAL)
    _, itf = textbox(s, 0.60, 1.42, 12.13, 0.4)
    para(itf, "Each feature is a guided page in the web app and is also exposed "
              "as MCP tools for AI clients.",
         size=12.5, color=GREY, first=True)

    feats = [
        ("Test Pilot (Advisor)", BLUE, BLUE_CHIP,
         "Chat with an AI agent to compare frameworks, get scored "
         "recommendations, and receive phased migration plans tailored to "
         "your project.",
         ["Scoring matrix", "Migration plan", "LLM chat"]),
        ("Framework Migration", GREEN, GREEN_CHIP,
         "Auto-convert test files between frameworks — single paste or full "
         "repo upload — preserving every assertion and validating the result.",
         ["Auto convert", "Multi-file", "Assertion-safe"]),
        ("Test Generator", ORANGE, ORANGE_CHIP,
         "Turn manual / plain-English test cases into runnable automated code "
         "with page objects, fixtures, and run commands.",
         ["Manual -> auto", "Page objects", "Multi-framework"]),
        ("Repo Comparator", PURPLE, PURPLE_CHIP,
         "Compare two test repos across different frameworks; match cases by "
         "name and report assertion-count parity so nothing is silently "
         "dropped.",
         ["Cross-repo", "Parity report", "ZIP upload"]),
        ("Coverage Analyzer", BLUE_DK, BLUE_CHIP,
         "Upload test cases (Excel / JSON / CSV) and see what percentage each "
         "framework actually covers, with a full matrix and gap details.",
         ["Excel upload", "Coverage matrix", "Best-fit rank"]),
        ("Team Profiles  (+)", RED, RED_CHIP,
         "Per-team coding standards, conversion rules, and reference docs so "
         "generated code matches how each team actually works.",
         ["Custom rules", "Reference docs", "Export / import"]),
    ]
    cw, ch, gx, gy = 3.94, 1.72, 0.17, 0.18
    x0, y0 = 0.60, 1.95
    for idx, (title, col, chip, desc, tags) in enumerate(feats):
        r, c = divmod(idx, 3)
        x = x0 + c * (cw + gx)
        y = y0 + r * (ch + gy)
        card = rounded(s, x, y, cw, ch, WHITE)
        card.line.color.rgb = chip
        card.line.width = Pt(1.25)
        rect(s, x, y, 0.09, ch, col)  # left accent spine
        _, ttf = textbox(s, x + 0.25, y + 0.13, cw - 0.4, 0.3)
        para(ttf, title, size=12.5, color=col, bold=True, first=True)
        _, dtf = textbox(s, x + 0.25, y + 0.48, cw - 0.45, 0.85)
        para(dtf, desc, size=9.3, color=GREY, first=True, line_spacing=1.03)
        _, gtf = textbox(s, x + 0.25, y + ch - 0.33, cw - 0.4, 0.28)
        gp = gtf.paragraphs[0]
        for ti, tag in enumerate(tags):
            run = gp.add_run()
            run.text = ("   " if ti else "") + "• " + tag
            _set_run(run, size=8.3, color=col, bold=True)

    _takeaway(s, "Five features cover the whole journey: choose -> migrate -> "
                 "generate -> verify, with per-team customization on top.")

    # ─────────────── SLIDE 5 — LLM + MCP SETUP / VALUE ─────────────────
    s = prs.slides.add_slide(blank)
    _page_chrome(s, "Connect an LLM & Configure via MCP", 5, TOTAL)

    # Left: LLM connection
    chip_header(s, 0.60, 1.45, 5.9, "1 — Connect an LLM provider", BLUE_CHIP, BLUE_DK)
    bullets(s, 0.60, 2.0, 5.9, 2.0, [
        "Groq or Gemini API key — added in the app's Settings, saved to "
        "config/.env, effective immediately (no restart).",
        "Or over MCP, use your AI client's own model via sampling — the "
        "server spends no key of its own.",
        "Falls back to GROQ_API_KEY automatically if the client can't sample.",
        "Non-AI features (scoring, coverage, parity) work with no key at all.",
    ], INK, size=10.8, gap=7)

    # code block: .env
    code = rounded(s, 0.60, 4.05, 5.9, 1.25, CODE_BG)
    code.line.fill.background()
    _, ctf = textbox(s, 0.78, 4.13, 5.6, 1.1)
    for i, ln in enumerate([
        "# config/.env",
        "GROQ_API_KEY = gsk_...",
        "GROQ_MODEL   = qwen/qwen3-32b",
        "# or",
        "GEMINI_API_KEY = AIza...",
    ]):
        para(ctf, ln, size=10, color=CODE_FG, font=MONO, first=(i == 0), space_after=2)

    chip_header(s, 0.60, 5.5, 5.9, "Works with: Groq  •  Gemini  •  client model (MCP)",
                GREEN_CHIP, GREEN)

    # Right: MCP configuration
    chip_header(s, 6.83, 1.45, 5.9, "2 — Configure as an MCP server", PURPLE_CHIP, PURPLE)
    bullets(s, 6.83, 2.0, 5.9, 1.7, [
        "Register the server once in mcp.json; the AI client launches it over "
        "stdio and all 20+ tools become callable from chat.",
        "DEFAULT_TEAM_PROFILE pins a team's rules for every call automatically.",
        "No UI needed — drive the whole advisor from your assistant.",
    ], INK, size=10.8, gap=7)

    code2 = rounded(s, 6.83, 3.75, 5.9, 1.95, CODE_BG)
    code2.line.fill.background()
    _, c2tf = textbox(s, 7.01, 3.83, 5.6, 1.8)
    for i, ln in enumerate([
        '"mcpServers": {',
        '  "automation-framework-advisor": {',
        '    "command": "python",',
        '    "args": ["mcp_server/server.py"],',
        '    "env": {',
        '      "DEFAULT_TEAM_PROFILE": "team_alpha"',
        '    }',
        '  }',
        '}',
    ]):
        para(c2tf, ln, size=9.5, color=CODE_FG, font=MONO, first=(i == 0), space_after=1)

    _takeaway(s, "Plug in one LLM key — or none, via MCP sampling — and the "
                 "whole advisor is usable from the web app or straight inside "
                 "your AI assistant.")

    out = Path(__file__).resolve().parent / "Test_Automation_Pilot.pptx"
    prs.save(out)
    print("Saved:", out)
    return out


if __name__ == "__main__":
    build()
